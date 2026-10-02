#!/usr/bin/env python3
"""Normalize FOIA workbooks and reconcile snapshots without modifying originals.

Requires Python 3.12+ and openpyxl. Dates remain Chicago local wall time.
"""

import argparse
import csv
import hashlib
import json
import shutil
from collections import Counter
from datetime import datetime
from pathlib import Path

import openpyxl


def workbook_rows(path: Path, sheet: str) -> list[dict[str, str]]:
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        iterator = book[sheet].iter_rows(values_only=True)
        headers = next(iterator)
        rows = []
        for values in iterator:
            if all(value is None for value in values):
                continue
            row = {}
            for key, value in zip(headers, values, strict=True):
                if isinstance(value, datetime):
                    row[key] = value.strftime("%m/%d/%Y %H:%M:%S")
                else:
                    row[key] = "" if value is None else str(value).strip()
            rows.append(row)
        return rows
    finally:
        book.close()


def write_csv(path: Path, rows: list[dict[str, str]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def index_tickets(rows: list[dict[str, str]]) -> dict[str, dict[str, str]]:
    indexed = {row["Ticket Number"]: row for row in rows}
    if len(indexed) != len(rows):
        raise ValueError("Duplicate ticket numbers: reconcile before building.")
    return indexed


def reconcile(old: list[dict[str, str]], new: list[dict[str, str]]) -> dict:
    old_index, new_index = index_tickets(old), index_tickets(new)
    shared = old_index.keys() & new_index.keys()
    changed = [ticket for ticket in sorted(shared) if old_index[ticket] != new_index[ticket]]
    changed_fields = Counter(
        key for ticket in changed for key in new_index[ticket]
        if old_index[ticket].get(key) != new_index[ticket][key]
    )
    old_end = max(datetime.strptime(row["Issued Date"], "%m/%d/%Y %H:%M:%S") for row in old)
    additions = [row for ticket, row in new_index.items() if ticket not in old_index]
    return {
        "previousRecords": len(old), "currentRecords": len(new),
        "retainedTickets": len(shared), "addedTickets": len(additions),
        "removedTickets": len(old_index.keys() - new_index.keys()),
        "changedTickets": len(changed), "changedFields": dict(changed_fields),
        "backfilledTickets": sum(datetime.strptime(row["Issued Date"], "%m/%d/%Y %H:%M:%S") <= old_end for row in additions),
        "previousLastIssuedAt": old_end.isoformat(),
        "previousListedFines": sum(float(row["Fine Level 1"] or 0) for row in old),
        "currentListedFines": sum(float(row["Fine Level 1"] or 0) for row in new),
        "changedTicketExamples": changed[:10],
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--smart-streets", type=Path, required=True)
    parser.add_argument("--illegal-parking", type=Path, required=True)
    parser.add_argument("--previous", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    smart = workbook_rows(args.smart_streets, "Sheet 1")
    parking = workbook_rows(args.illegal_parking, "Data")
    index_tickets(parking)
    for path, rows in [(args.smart_streets, smart), (args.illegal_parking, parking)]:
        shutil.copyfile(path, args.output_dir / path.name)
        write_csv(args.output_dir / path.with_suffix(".csv").name, rows)
    header_book = openpyxl.load_workbook(args.illegal_parking, read_only=True, data_only=True)
    notes = [str(value) for row in header_book["Header"].iter_rows(values_only=True) for value in row if value is not None]
    header_book.close()
    report = {
        "smartStreets": reconcile(read_csv(args.previous), smart),
        "illegalParking": {
            "records": len(parking),
            "missingLocationRecords": sum(row["Location"] == "" for row in parking),
            "violationDescriptions": dict(Counter(row["Violation Description"] for row in parking)),
            "violationCodes": dict(Counter(row["Violation Code"] for row in parking)),
            "originalNotes": notes,
        },
        "originalWorkbooks": [{"filename": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()} for path in [args.smart_streets, args.illegal_parking]],
    }
    (args.output_dir / "foia-reconciliation.json").write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
